from __future__ import annotations

import csv
import html
import io
import json
from dataclasses import asdict, dataclass
from datetime import datetime
from typing import Literal

from openpyxl import Workbook  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.styles import getSampleStyleSheet  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore[import-untyped]

ReportKind = Literal["summary", "detailed"]


@dataclass(frozen=True, slots=True)
class ReportEntry:
    profile: str
    connector: str
    url: str
    title: str
    snippet: str
    confidence: float
    verification_status: str
    discovered_at: str
    last_checked_at: str
    published_at: str | None


@dataclass(frozen=True, slots=True)
class ReportData:
    generated_at: datetime
    scope: str
    sources_queried: tuple[str, ...]
    sources_unavailable: tuple[str, ...]
    entries: tuple[ReportEntry, ...]
    methodological_limits: str = (
        "Les résultats publics peuvent être incomplets ou obsolètes. "
        "Un score est un indice et ne prouve jamais l'identité."
    )


HEADERS = (
    "Profil",
    "Source",
    "URL",
    "Titre",
    "Extrait",
    "Confiance",
    "Validation",
    "Découverte",
    "Dernière vérification",
    "Publication",
)


def neutralize_spreadsheet(value: object) -> str:
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text


def _rows(report: ReportData) -> list[tuple[object, ...]]:
    return [
        (
            item.profile,
            item.connector,
            item.url,
            item.title,
            item.snippet,
            item.confidence,
            item.verification_status,
            item.discovered_at,
            item.last_checked_at,
            item.published_at,
        )
        for item in report.entries
    ]


def render_csv(report: ReportData, kind: ReportKind) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(("Périmètre", neutralize_spreadsheet(report.scope)))
    writer.writerow(("Généré le", report.generated_at.isoformat()))
    writer.writerow(("Sources interrogées", ", ".join(report.sources_queried)))
    writer.writerow(("Sources indisponibles", ", ".join(report.sources_unavailable)))
    writer.writerow(("Nombre de résultats", len(report.entries)))
    writer.writerow(("Limites", neutralize_spreadsheet(report.methodological_limits)))
    if kind == "detailed":
        writer.writerow(HEADERS)
        writer.writerows(
            tuple(neutralize_spreadsheet(value) for value in row) for row in _rows(report)
        )
    return output.getvalue().encode("utf-8-sig")


def render_json(report: ReportData, kind: ReportKind) -> bytes:
    payload = {
        "generated_at": report.generated_at.isoformat(),
        "scope": report.scope,
        "sources_queried": report.sources_queried,
        "sources_unavailable": report.sources_unavailable,
        "result_count": len(report.entries),
        "methodological_limits": report.methodological_limits,
        "results": [asdict(entry) for entry in report.entries] if kind == "detailed" else [],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode()


def render_xlsx(report: ReportData, kind: ReportKind) -> bytes:
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Synthèse")
    summary.append(("Périmètre", neutralize_spreadsheet(report.scope)))
    summary.append(("Généré le", report.generated_at.isoformat()))
    summary.append(("Sources interrogées", ", ".join(report.sources_queried)))
    summary.append(("Sources indisponibles", ", ".join(report.sources_unavailable)))
    summary.append(("Nombre de résultats", len(report.entries)))
    summary.append(("Limites", neutralize_spreadsheet(report.methodological_limits)))
    if kind == "detailed":
        details = workbook.create_sheet("Mentions")
        details.append(HEADERS)
        for row in _rows(report):
            details.append(tuple(neutralize_spreadsheet(value) for value in row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_html(report: ReportData, kind: ReportKind) -> bytes:
    def escaped(value: object) -> str:
        return html.escape("" if value is None else str(value), quote=True)

    rows = ""
    if kind == "detailed":
        rows = "".join(
            "<tr>" + "".join(f"<td>{escaped(value)}</td>" for value in row) + "</tr>"
            for row in _rows(report)
        )
    headings = "".join(f"<th>{escaped(header)}</th>" for header in HEADERS)
    styles = (
        "body{font-family:sans-serif;margin:2rem}"
        "table{border-collapse:collapse;width:100%}"
        "th,td{border:1px solid #bbb;padding:.4rem;text-align:left;vertical-align:top}"
        "th{background:#eee}"
    )
    document = "".join(
        (
            '<!doctype html><html lang="fr"><head><meta charset="utf-8">',
            '<meta http-equiv="Content-Security-Policy" ',
            'content="default-src \'none\'; style-src \'unsafe-inline\'">',
            f"<title>Rapport DigitalFootprint</title><style>{styles}</style></head>",
            f"<body><h1>Rapport DigitalFootprint</h1><p><strong>Périmètre :</strong> "
            f"{escaped(report.scope)}</p>",
            f"<p><strong>Généré le :</strong> {escaped(report.generated_at.isoformat())}</p>",
            f"<p><strong>Sources interrogées :</strong> "
            f"{escaped(', '.join(report.sources_queried))}</p>",
            f"<p><strong>Sources indisponibles :</strong> "
            f"{escaped(', '.join(report.sources_unavailable))}</p>",
            f"<p><strong>Nombre de résultats :</strong> {len(report.entries)}</p>",
            f"<p><strong>Limites :</strong> {escaped(report.methodological_limits)}</p>",
            f"<table><thead><tr>{headings}</tr></thead><tbody>{rows}</tbody></table>",
            "</body></html>",
        )
    )
    return document.encode()


def render_pdf(report: ReportData, kind: ReportKind) -> bytes:
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Rapport DigitalFootprint",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Rapport DigitalFootprint", styles["Title"]),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>Périmètre :</b> {html.escape(report.scope)}", styles["BodyText"]),
        Paragraph(f"<b>Résultats :</b> {len(report.entries)}", styles["BodyText"]),
        Paragraph(
            f"<b>Sources :</b> {html.escape(', '.join(report.sources_queried))}",
            styles["BodyText"],
        ),
        Paragraph(
            f"<b>Limites :</b> {html.escape(report.methodological_limits)}",
            styles["BodyText"],
        ),
    ]
    if kind == "detailed":
        for entry in report.entries:
            story.extend(
                [
                    Spacer(1, 4 * mm),
                    Paragraph(html.escape(entry.title), styles["Heading3"]),
                    Paragraph(html.escape(entry.url), styles["BodyText"]),
                    Paragraph(html.escape(entry.snippet), styles["BodyText"]),
                    Paragraph(
                        f"Confiance : {entry.confidence:.0%} — "
                        f"{html.escape(entry.verification_status)}",
                        styles["BodyText"],
                    ),
                ]
            )
    document.build(story)
    return output.getvalue()
